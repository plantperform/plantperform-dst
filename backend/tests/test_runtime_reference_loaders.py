import csv
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import openpyxl

from database.scripts import runtime_lookup_loader
from database.scripts.load_afgroede_normer import parse_afgroede_normer
from database.scripts.load_afstromningskategorier import parse_afstromningskategorier
from database.scripts.load_arbejdsmaengder import parse_arbejdsmaengder
from database.scripts.load_arbejdssatser import parse_arbejdssatser
from database.scripts.load_dyrkningsomkostninger import parse_dyrkningsomkostninger
from database.scripts.load_halmudbytte import parse_halmudbytte
from database.scripts.load_prisliste import parse_prisliste
from database.scripts.load_salgspriser import parse_salgspriser


def write_csv(path: Path, headers: list[str], row: list[str], *, delimiter: str = ",") -> None:
    with path.open("w", encoding="utf-8", newline="") as output:
        writer = csv.writer(output, delimiter=delimiter)
        writer.writerow(headers)
        writer.writerow(row)


class RuntimeReferenceLoaderTests(unittest.TestCase):
    def test_each_csv_loader_parses_its_own_source(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            base = Path(directory)
            runoff = base / "runoff.csv"
            write_csv(
                runoff,
                ["Afgrødekode", "P_afstrømningskategori", "P_afstrømningskategori_med_W"],
                ["1", "3", "1"],
                delimiter=";",
            )
            sales = base / "sales.csv"
            write_csv(
                sales,
                ["afgroedekode", "driftsform", "kvalitet", "salgspris", "enhed", "halm_pris_kr_kg"],
                ["1", "Konventionel", "", "1.5", "kg", ""],
            )
            straw = base / "straw.csv"
            write_csv(
                straw,
                ["afgroedekode", "jordbonitet", "halm_udbytte_kg_ha"],
                ["1", "JB1-3", "2"],
            )
            rates = base / "rates.csv"
            write_csv(
                rates,
                ["behandling", "jordbonitet", "afgroedekode", "driftsform", "pris_kr_per_enhed"],
                ["Såning", "JB1-3", "", "", "3"],
            )
            quantities = base / "quantities.csv"
            write_csv(
                quantities,
                [
                    "afgroedekode", "driftsform", "jordbonitet", "kvalitet", "kategori",
                    "behandling", "antal",
                ],
                ["1", "Konventionel", "JB1-3", "", "Udsæd", "Såning", "2"],
            )
            costs = base / "costs.csv"
            write_csv(
                costs,
                ["afgroedekode", "driftsform", "kategori", "behandling", "udgift_kr_ha"],
                ["1", "Konventionel", "Udsæd", "Korn", "12"],
            )
            prices = base / "prices.csv"
            write_csv(
                prices,
                ["post", "kategori", "type", "pris", "enhed"],
                ["N", "Gødning", "Omkostning", "4", "kr/kg"],
                delimiter=";",
            )

            self.assertEqual(parse_afstromningskategorier(runoff), [(1, 3, 1)])
            self.assertEqual(parse_salgspriser(sales)[0][-2:], ("kg", 0.0))
            self.assertEqual(parse_halmudbytte(straw)[0][-1], 2.0)
            self.assertIsNone(parse_arbejdssatser(rates)[0][3])
            self.assertEqual(parse_arbejdsmaengder(quantities)[0][-1], 2.0)
            self.assertEqual(parse_dyrkningsomkostninger(costs)[0][-1], 12.0)
            self.assertEqual(parse_prisliste(prices)[0][-1], "kr/kg")

    def test_master_workbook_loader_requires_and_parses_all_three_sheets(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "normer.xlsx"
            workbook = openpyxl.Workbook()
            lang = workbook.active
            lang.title = "Lang_lookup"
            lang.append([
                "Afgrødekode", "Afgrøde", "JB_gruppe", "JB_match_type", "JB_værdier",
                "Vanding", "Udbytteenhed", "Udbyttenorm", "Udbyttenorm_alt_ikke_korn",
                "N_norm_kgN_ha", "P_norm_kgP_ha", "Forfrugtsværdi_kgN_ha",
                "Indregn_forfrugtsværdi_i_N_norm", "Driftsform",
            ])
            lang.append([
                1, "Vårbyg", "JB 1", "enkelt", "1", "Uvandet", "hkg", 50, None, 100, 20,
                5, "Ja", None,
            ])
            nfix = workbook.create_sheet("N_fixering_lookup")
            nfix.append(["Afgrødekode", "JB_værdier", "Vanding", "Nfix_kgN_ha"])
            nfix.append([1, "1", "Uvandet", 7])
            nuar = workbook.create_sheet("NUAR_koder")
            nuar.append([
                "AfgroedeKode", "Navn", "M", "W", "WC", "MP", "WP", "M_ambig", "W_ambig",
                "WC_ambig", "MP_ambig", "WP_ambig",
            ])
            nuar.append([1, "Vårbyg", 1, 2, 3, 4, 5, 0, 1, 0, 1, 0])
            workbook.save(path)
            workbook.close()

            norm_rows, nfix_rows, nuar_rows = parse_afgroede_normer(path)

        self.assertEqual(norm_rows[0][2], 1)
        self.assertEqual(nfix_rows, [(2, 1, 1, "Uvandet", 7.0)])
        self.assertEqual(nuar_rows[0][0:3], (1, "Vårbyg", 1))

    def test_empty_replacement_never_opens_a_database_connection(self) -> None:
        with patch.object(runtime_lookup_loader.psycopg, "connect") as connect:
            with self.assertRaisesRegex(ValueError, "no rows"):
                runtime_lookup_loader.replace_tables([("salgspris", ("source_order",), [])])
        connect.assert_not_called()


if __name__ == "__main__":
    unittest.main()
