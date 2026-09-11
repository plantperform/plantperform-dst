"""Load crop norms, N fixation, and NUAR codes from the master workbook."""

from __future__ import annotations

from collections.abc import Iterable
from pathlib import Path

import openpyxl

from database.scripts.runtime_lookup_loader import replace_tables, source_path

XLSX_PATH = source_path(
    "PlantPerform_master_afgroedenormer_opdateret_fra_hoeringsmateriale_Bilag_1_1_2027.xlsx",
)
KONVENTIONEL = "Konventionel"


def _text(value: object) -> str:
    return str(value).strip() if value is not None else ""


def _number(value: object) -> float | None:
    if value is None:
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        raw = _text(value)
        if "/" in raw:
            try:
                return float(raw.split("/", maxsplit=1)[0])
            except ValueError:
                pass
    return None


def _integer(value: object, *, field: str, row_number: int) -> int:
    parsed = _number(value)
    if parsed is None or not parsed.is_integer():
        raise ValueError(f"Invalid {field} {value!r} on row {row_number}")
    return int(parsed)


def _optional_integer(value: object) -> int | None:
    parsed = _number(value)
    return int(parsed) if parsed is not None and parsed.is_integer() else None


def _jb_values(match_type: object, values: object, *, row_number: int) -> Iterable[int]:
    raw = _text(values)
    if not raw or raw.lower() == "none":
        raise ValueError(f"Missing JB_værdier on row {row_number}")
    kind = _text(match_type).lower()
    try:
        if kind == "plus":
            return sorted({int(value) for value in raw.split(";")})
        if kind == "til":
            first, last = raw.split("-", maxsplit=1)
            return range(int(first), int(last) + 1)
        if kind == "enkelt":
            return [int(raw)]
    except ValueError as error:
        raise ValueError(f"Invalid JB_værdier {raw!r} on row {row_number}") from error
    raise ValueError(f"Invalid JB_match_type {kind!r} on row {row_number}")


def _jb_values_inferred(values: object, *, row_number: int) -> Iterable[int]:
    raw = _text(values)
    if not raw or raw.lower() == "none":
        raise ValueError(f"Missing JB_værdier on row {row_number}")
    try:
        if ";" in raw:
            return sorted({int(value) for value in raw.split(";")})
        if "-" in raw:
            first, last = raw.split("-", maxsplit=1)
            return range(int(first), int(last) + 1)
        return [int(raw)]
    except ValueError as error:
        raise ValueError(f"Invalid JB_værdier {raw!r} on row {row_number}") from error


def _sheet_rows(workbook, sheet_name: str, required_headers: set[str]):
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Workbook is missing sheet {sheet_name!r}")
    sheet = workbook[sheet_name]
    headers = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    by_name = {_text(header): index for index, header in enumerate(headers) if _text(header)}
    missing = required_headers - set(by_name)
    if missing:
        raise ValueError(f"{sheet_name} is missing columns: {', '.join(sorted(missing))}")
    return sheet, by_name


def _cell(row: tuple, columns: dict[str, int], name: str) -> object:
    index = columns[name]
    return row[index] if index < len(row) else None


def parse_afgroede_normer(path: Path = XLSX_PATH) -> tuple[list[tuple], list[tuple], list[tuple]]:
    """Parse all workbook sheets required by crop calculations before DB writes."""
    if not path.exists():
        raise FileNotFoundError(f"Expected {path}")
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        lang_sheet, lang_columns = _sheet_rows(
            workbook,
            "Lang_lookup",
            {
                "Afgrødekode", "Afgrøde", "JB_gruppe", "JB_match_type", "JB_værdier",
                "Vanding", "Udbytteenhed", "Udbyttenorm", "Udbyttenorm_alt_ikke_korn",
                "N_norm_kgN_ha", "P_norm_kgP_ha", "Forfrugtsværdi_kgN_ha",
                "Indregn_forfrugtsværdi_i_N_norm", "Driftsform",
            },
        )
        norm_rows: list[tuple] = []
        for source_order, row in enumerate(
            lang_sheet.iter_rows(min_row=2, values_only=True), start=2,
        ):
            if not _text(_cell(row, lang_columns, "Afgrødekode")):
                continue
            crop_code = _integer(
                _cell(row, lang_columns, "Afgrødekode"),
                field="Afgrødekode",
                row_number=source_order,
            )
            values = _jb_values(
                _cell(row, lang_columns, "JB_match_type"),
                _cell(row, lang_columns, "JB_værdier"),
                row_number=source_order,
            )
            common = (
                crop_code,
                _text(_cell(row, lang_columns, "Afgrøde")),
                _text(_cell(row, lang_columns, "JB_gruppe")),
                _text(_cell(row, lang_columns, "Vanding")),
                _text(_cell(row, lang_columns, "Udbytteenhed")),
                _number(_cell(row, lang_columns, "Udbyttenorm")),
                _number(_cell(row, lang_columns, "Udbyttenorm_alt_ikke_korn")),
                _number(_cell(row, lang_columns, "N_norm_kgN_ha")),
                _number(_cell(row, lang_columns, "P_norm_kgP_ha")),
                _number(_cell(row, lang_columns, "Forfrugtsværdi_kgN_ha")) or 0.0,
                _text(_cell(row, lang_columns, "Indregn_forfrugtsværdi_i_N_norm")).lower()
                == "ja",
                _text(_cell(row, lang_columns, "Driftsform")) or KONVENTIONEL,
            )
            norm_rows.extend((source_order, jb_nr, *common) for jb_nr in values)

        nfix_sheet, nfix_columns = _sheet_rows(
            workbook,
            "N_fixering_lookup",
            {"Afgrødekode", "JB_værdier", "Vanding", "Nfix_kgN_ha"},
        )
        nfix_rows: list[tuple] = []
        for source_order, row in enumerate(
            nfix_sheet.iter_rows(min_row=2, values_only=True), start=2,
        ):
            if not _text(_cell(row, nfix_columns, "Afgrødekode")):
                continue
            nfix = _number(_cell(row, nfix_columns, "Nfix_kgN_ha"))
            if nfix is None:
                continue
            crop_code = _integer(
                _cell(row, nfix_columns, "Afgrødekode"),
                field="Afgrødekode",
                row_number=source_order,
            )
            nfix_rows.extend(
                (source_order, jb_nr, crop_code, _text(_cell(row, nfix_columns, "Vanding")), nfix)
                for jb_nr in _jb_values_inferred(
                    _cell(row, nfix_columns, "JB_værdier"), row_number=source_order,
                )
            )

        nuar_sheet, nuar_columns = _sheet_rows(
            workbook,
            "NUAR_koder",
            {
                "AfgroedeKode", "Navn", "M", "W", "WC", "MP", "WP", "M_ambig",
                "W_ambig", "WC_ambig", "MP_ambig", "WP_ambig",
            },
        )
        nuar_rows: list[tuple] = []
        for row_number, row in enumerate(
            nuar_sheet.iter_rows(min_row=2, values_only=True), start=2,
        ):
            if not _text(_cell(row, nuar_columns, "AfgroedeKode")):
                continue
            nuar_rows.append(
                (
                    _integer(
                        _cell(row, nuar_columns, "AfgroedeKode"),
                        field="AfgroedeKode",
                        row_number=row_number,
                    ),
                    _text(_cell(row, nuar_columns, "Navn")),
                    _optional_integer(_cell(row, nuar_columns, "M")),
                    _optional_integer(_cell(row, nuar_columns, "W")),
                    _optional_integer(_cell(row, nuar_columns, "WC")),
                    _optional_integer(_cell(row, nuar_columns, "MP")),
                    _optional_integer(_cell(row, nuar_columns, "WP")),
                    bool(_cell(row, nuar_columns, "M_ambig")),
                    bool(_cell(row, nuar_columns, "W_ambig")),
                    bool(_cell(row, nuar_columns, "WC_ambig")),
                    bool(_cell(row, nuar_columns, "MP_ambig")),
                    bool(_cell(row, nuar_columns, "WP_ambig")),
                )
            )
    finally:
        workbook.close()

    if not norm_rows or not nfix_rows or not nuar_rows:
        raise ValueError("Master workbook must contain norm, N fixation, and NUAR data")
    if len({row[0] for row in nuar_rows}) != len(nuar_rows):
        raise ValueError("NUAR_koder contains duplicate afgrødekoder")
    return norm_rows, nfix_rows, nuar_rows


def load_afgroede_normer(path: Path = XLSX_PATH, database_url: str | None = None) -> None:
    norm_rows, nfix_rows, nuar_rows = parse_afgroede_normer(path)
    replace_tables(
        [
            (
                "afgroede_norm_lookup",
                (
                    "source_order", "jb_nr", "afgroedekode", "afgroede", "jb_gruppe", "vanding",
                    "udbytteenhed", "udbyttenorm", "udbyttenorm_alt", "n_norm", "p_norm",
                    "forfrugtsvaerdi", "indregn_ffv", "driftsform",
                ),
                norm_rows,
            ),
            (
                "afgroede_nfix_lookup",
                ("source_order", "jb_nr", "afgroedekode", "vanding", "nfix_kgn_ha"),
                nfix_rows,
            ),
            (
                "nuar_kode",
                (
                    "afgroedekode", "navn", "m", "w", "wc", "mp", "wp", "m_ambig",
                    "w_ambig", "wc_ambig", "mp_ambig", "wp_ambig",
                ),
                nuar_rows,
            ),
        ],
        database_url=database_url,
    )
    print(
        f"Loaded {len(norm_rows):,} crop-norm rows, {len(nfix_rows):,} N-fixation rows, "
        f"and {len(nuar_rows):,} NUAR codes",
        flush=True,
    )


if __name__ == "__main__":
    load_afgroede_normer()
